#!/usr/bin/env python3
"""Shared rules for Tea Card Salt Lake daily tourism-weather indices.

The historical baseline builder and the future-forecast calculator are separate
entry points. This module contains the unchanged formulas, quality controls,
baseline serialization, and forecast calculation shared by both programs.
"""

from __future__ import annotations

import bisect
import hashlib
import json
import math
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl


STATION_LATITUDE_DEG = 36.7833
STATION_ELEVATION_M = 3087.6
METHOD_VERSION = "1.1-altitude-uvp"
BASELINE_SCHEMA_VERSION = "1.0"
MISSING_SENTINEL = 999000.0
SUNSHINE_TOLERANCE_HOURS = 0.05
UVP_REFERENCE_MAX = 14.131644170491846
UVP_ALTITUDE_INCREASE_PER_1000M = 0.10

REQUIRED_HISTORY_COLUMNS = {
    "year": "年", "month": "月", "day": "日", "avg_temp": "平均气温",
    "min_temp": "最低气温", "avg_rh": "平均相对湿度",
    "avg_wind": "平均2分钟风速", "sunshine_hours": "日照时数（直接辐射计算值）",
}


class InputError(ValueError):
    """Raised when an input cannot be used by the agreed rules."""


@dataclass(frozen=True)
class Baseline:
    min_temp_values: tuple[float, ...]
    k_values: tuple[float, ...]
    uvp_thresholds: tuple[float, float, float, float]
    clothing_thresholds: tuple[float, float, float, float]
    reference_start: str
    reference_end: str
    valid_counts: dict[str, int]
    station_elevation_m: float
    uvp_altitude_factor: float
    source_workbook_name: str
    source_workbook_sha256: str


def _number(value: Any) -> float | None:
    """Coerce a number; any 999xxx sentinel is missing."""
    if value is None:
        return None
    try:
        value = float(value)
    except (TypeError, ValueError):
        return None
    return value if math.isfinite(value) and value < MISSING_SENTINEL else None


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _valid(value: float | None, lower: float, upper: float) -> bool:
    return value is not None and lower <= value <= upper


def _parse_date(value: str | date) -> date:
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except ValueError as exc:
        raise InputError("date must use YYYY-MM-DD, for example 2026-07-20") from exc


def solar_geometry(target_date: date, latitude_deg: float = STATION_LATITUDE_DEG) -> tuple[float, float]:
    """Return theoretical day length (hours) and sin(noon solar altitude)."""
    d = target_date.timetuple().tm_yday
    phi = math.radians(latitude_deg)
    declination = math.radians(23.45) * math.sin(2 * math.pi * (284 + d) / 365.0)
    cos_hour_angle = max(-1.0, min(1.0, -math.tan(phi) * math.tan(declination)))
    day_length = 2 * math.degrees(math.acos(cos_hour_angle)) / 15.0
    noon_altitude = math.asin(math.sin(phi) * math.sin(declination) + math.cos(phi) * math.cos(declination))
    return day_length, max(0.0, math.sin(noon_altitude))


def normalise_sunshine(target_date: date, sunshine_hours: float) -> float:
    """Apply the physical day-length quality-control rule."""
    if not _valid(sunshine_hours, 0.0, 24.0):
        raise InputError("sunshine_hours must be between 0 and 24 and cannot be a 999xxx missing code")
    day_length, _ = solar_geometry(target_date)
    if sunshine_hours > day_length + SUNSHINE_TOLERANCE_HOURS:
        raise InputError(
            f"sunshine_hours ({sunshine_hours}) exceeds theoretical day length ({day_length:.2f}) "
            f"by more than {SUNSHINE_TOLERANCE_HOURS:.2f} hours"
        )
    return min(sunshine_hours, day_length)


def calculate_thi(avg_temp: float, avg_rh: float) -> float:
    f = avg_rh / 100.0
    return (1.8 * avg_temp + 32) - 0.55 * (1 - f) * (1.8 * avg_temp - 26)


def calculate_k(avg_temp: float, avg_wind: float, sunshine_hours: float) -> float:
    return -(10 * math.sqrt(avg_wind) + 10.45 - avg_wind) * (33 - avg_temp) + 8.55 * sunshine_hours


def uvp_altitude_factor(elevation_m: float = STATION_ELEVATION_M) -> float:
    """WHO approximate UV increase: 10% for each 1,000 m of elevation."""
    return 1.0 + UVP_ALTITUDE_INCREASE_PER_1000M * elevation_m / 1000.0


def calculate_uvp(target_date: date, sunshine_hours: float) -> tuple[float, float]:
    """Return the unadjusted and altitude-adjusted local UVP values."""
    _, noon_sine = solar_geometry(target_date)
    unadjusted = 100.0 * sunshine_hours * noon_sine / UVP_REFERENCE_MAX
    return unadjusted, unadjusted * uvp_altitude_factor()


def thi_score_and_level(thi: float) -> tuple[int, str]:
    if thi < 40: return -8, "极冷"
    if thi < 45: return -6, "寒冷"
    if thi < 55: return -4, "偏冷"
    if thi < 60: return -2, "偏凉"
    if thi < 65: return 0, "中性"
    if thi < 70: return 2, "偏暖"
    if thi < 75: return 4, "暖热"
    if thi < 80: return 6, "闷热"
    return 8, "极闷热"


def k_score_and_level(k: float) -> tuple[int, str]:
    if k < -1200: return -8, "酷冷"
    if k < -1000: return -6, "冷"
    if k < -800: return -4, "冷凉"
    if k < -600: return -2, "凉"
    if k < -300: return 0, "舒适"
    if k < -200: return 2, "暖"
    if k < -50: return 4, "暖热"
    if k < 80: return 6, "热"
    return 8, "炎热"


def _quantile(values: list[float], probability: float) -> float:
    if not values:
        raise ValueError("reference distribution is empty")
    position = (len(values) - 1) * probability
    lower, upper = math.floor(position), math.ceil(position)
    return values[lower] + (values[upper] - values[lower]) * (position - lower)


def _cold_rank(reference: tuple[float, ...], value: float) -> float:
    """Return 0-100; lower temperature/K has a higher cold rank."""
    return 100.0 * (len(reference) - bisect.bisect_left(reference, value)) / len(reference)


def _five_level(value: float, thresholds: tuple[float, float, float, float], labels: tuple[str, ...]) -> str:
    for threshold, label in zip(thresholds, labels):
        if value < threshold:
            return label
    return labels[-1]


def _validate_forecast(forecast: dict[str, Any]) -> tuple[date, float, float, float, float, float]:
    fields = ("date", "avg_temp", "min_temp", "avg_rh", "avg_wind", "sunshine_hours")
    absent = [field for field in fields if field not in forecast]
    if absent:
        raise InputError(f"missing required fields: {', '.join(absent)}")
    target_date = _parse_date(forecast["date"])
    values = {field: _number(forecast[field]) for field in fields if field != "date"}
    if not _valid(values["avg_temp"], -80, 70):
        raise InputError("avg_temp must be between -80 and 70 °C")
    if not _valid(values["min_temp"], -80, 70):
        raise InputError("min_temp must be between -80 and 70 °C")
    if not _valid(values["avg_rh"], 0, 100):
        raise InputError("avg_rh must be between 0 and 100 percent")
    if not _valid(values["avg_wind"], 0, 100):
        raise InputError("avg_wind must be between 0 and 100 m/s")
    sunshine = normalise_sunshine(target_date, values["sunshine_hours"])
    return target_date, values["avg_temp"], values["min_temp"], values["avg_rh"], values["avg_wind"], sunshine


def build_baseline(history_workbook: str | Path) -> Baseline:
    """Build historical reference distributions from a Tea Card station workbook."""
    history_path = Path(history_workbook)
    workbook = openpyxl.load_workbook(history_path, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    positions = {header: idx for idx, header in enumerate(headers)}
    missing = [column for column in REQUIRED_HISTORY_COLUMNS.values() if column not in positions]
    if missing:
        raise InputError(f"history workbook is missing columns: {', '.join(missing)}")

    min_temps: list[float] = []
    k_values: list[float] = []
    uvps: list[float] = []
    paired: list[tuple[float, float]] = []
    dates: list[date] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        try:
            row_date = date(
                int(row[positions[REQUIRED_HISTORY_COLUMNS["year"]]]),
                int(row[positions[REQUIRED_HISTORY_COLUMNS["month"]]]),
                int(row[positions[REQUIRED_HISTORY_COLUMNS["day"]]]),
            )
        except (TypeError, ValueError):
            continue
        dates.append(row_date)
        t = _number(row[positions[REQUIRED_HISTORY_COLUMNS["avg_temp"]]])
        tmin = _number(row[positions[REQUIRED_HISTORY_COLUMNS["min_temp"]]])
        wind = _number(row[positions[REQUIRED_HISTORY_COLUMNS["avg_wind"]]])
        sunshine = _number(row[positions[REQUIRED_HISTORY_COLUMNS["sunshine_hours"]]])
        if _valid(tmin, -80, 70):
            min_temps.append(tmin)
        try:
            sunshine = normalise_sunshine(row_date, sunshine) if sunshine is not None else None
        except InputError:
            sunshine = None
        if sunshine is not None:
            _, uvp = calculate_uvp(row_date, sunshine)
            uvps.append(uvp)
        if _valid(t, -80, 70) and _valid(wind, 0, 100) and sunshine is not None:
            k = calculate_k(t, wind, sunshine)
            k_values.append(k)
            if _valid(tmin, -80, 70):
                paired.append((tmin, k))
    if not dates or not min_temps or not k_values or not uvps or not paired:
        raise InputError("history workbook does not contain enough valid records")

    min_temps.sort()
    k_values.sort()
    min_temp_ref, k_ref = tuple(min_temps), tuple(k_values)
    clothing = sorted(max(_cold_rank(min_temp_ref, tmin), _cold_rank(k_ref, k)) for tmin, k in paired)
    uvps.sort()
    return Baseline(
        min_temp_values=min_temp_ref,
        k_values=k_ref,
        uvp_thresholds=tuple(_quantile(uvps, q) for q in (0.2, 0.4, 0.6, 0.8)),
        clothing_thresholds=tuple(_quantile(clothing, q) for q in (0.2, 0.4, 0.6, 0.8)),
        reference_start=min(dates).isoformat(),
        reference_end=max(dates).isoformat(),
        valid_counts={"min_temp": len(min_temps), "wind_effect_k": len(k_values), "uvp": len(uvps), "clothing": len(clothing)},
        station_elevation_m=STATION_ELEVATION_M,
        uvp_altitude_factor=uvp_altitude_factor(),
        source_workbook_name=history_path.name,
        source_workbook_sha256=_sha256_file(history_path),
    )


def _style_baseline_sheet(worksheet: Any, widths: tuple[float, ...]) -> None:
    """Apply light audit-oriented formatting to one baseline workbook sheet."""
    header_fill = openpyxl.styles.PatternFill("solid", fgColor="1F4E78")
    header_font = openpyxl.styles.Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width


def write_baseline_xlsx(baseline: Baseline, output_path: str | Path) -> None:
    """Write a frozen, auditable baseline workbook for the forecast calculator."""
    output = Path(output_path)
    if output.suffix.lower() != ".xlsx":
        raise InputError("baseline output must use the .xlsx extension")
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.Workbook()
    metadata = workbook.active
    metadata.title = "Metadata"
    metadata.append(["field", "value", "description"])
    metadata_rows = [
        ("baseline_schema_version", BASELINE_SCHEMA_VERSION, "Baseline workbook structure version"),
        ("method_version", METHOD_VERSION, "Index method version"),
        ("source_workbook_name", baseline.source_workbook_name, "Historical workbook used to construct the baseline"),
        ("source_workbook_sha256", baseline.source_workbook_sha256, "SHA-256 of the historical workbook"),
        ("reference_start", baseline.reference_start, "First valid date parsed from historical workbook"),
        ("reference_end", baseline.reference_end, "Last valid date parsed from historical workbook"),
        ("station_latitude_deg", STATION_LATITUDE_DEG, "Fixed latitude used for solar geometry"),
        ("station_elevation_m", baseline.station_elevation_m, "Fixed elevation used only for UVP"),
        ("uvp_altitude_factor", baseline.uvp_altitude_factor, "UVP multiplier applied to history and forecast"),
        ("missing_sentinel", MISSING_SENTINEL, "Values greater than or equal to this are missing"),
        ("sunshine_tolerance_hours", SUNSHINE_TOLERANCE_HOURS, "Allowed sunshine excess over theoretical day length"),
        ("uvp_reference_max", UVP_REFERENCE_MAX, "UVP solar-geometry normalization constant"),
    ]
    for row in metadata_rows:
        metadata.append(row)
    _style_baseline_sheet(metadata, (32, 26, 62))

    thresholds = workbook.create_sheet("Thresholds")
    thresholds.append(["index", "quantile", "threshold_value", "meaning"])
    for index_name, values, meaning in (
        ("uvp", baseline.uvp_thresholds, "Historical UVP cut point"),
        ("clothing", baseline.clothing_thresholds, "Historical clothing-demand cut point"),
    ):
        for quantile, value in zip((0.2, 0.4, 0.6, 0.8), values):
            thresholds.append([index_name, quantile, value, meaning])
    _style_baseline_sheet(thresholds, (18, 14, 22, 42))
    for row in thresholds.iter_rows(min_row=2, max_row=9, min_col=2, max_col=3):
        row[0].number_format = "0%"
        row[1].number_format = "0.000"

    distributions = workbook.create_sheet("Distributions")
    distributions.append(["rank", "min_temp_c_sorted", "wind_effect_k_sorted"])
    distribution_length = max(len(baseline.min_temp_values), len(baseline.k_values))
    for rank in range(distribution_length):
        distributions.append([
            rank + 1,
            baseline.min_temp_values[rank] if rank < len(baseline.min_temp_values) else None,
            baseline.k_values[rank] if rank < len(baseline.k_values) else None,
        ])
    _style_baseline_sheet(distributions, (12, 22, 24))
    for row in distributions.iter_rows(min_row=2, max_row=distribution_length + 1, min_col=2, max_col=3):
        for cell in row:
            cell.number_format = "0.000"

    validation = workbook.create_sheet("Validation")
    validation.append(["metric", "valid_days", "description"])
    descriptions = {
        "min_temp": "Valid minimum-temperature records",
        "wind_effect_k": "Valid K records after sunshine quality control",
        "uvp": "Valid UVP records after sunshine quality control",
        "clothing": "Records with both valid Tmin and K",
    }
    for metric, count in baseline.valid_counts.items():
        validation.append([metric, count, descriptions.get(metric, "")])
    _style_baseline_sheet(validation, (22, 16, 56))

    workbook.save(output)
    workbook.close()


def _metadata_value(metadata: dict[str, Any], key: str) -> Any:
    if key not in metadata or metadata[key] is None:
        raise InputError(f"baseline workbook metadata is missing '{key}'")
    return metadata[key]


def _as_float(value: Any, field: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise InputError(f"baseline workbook field '{field}' must be numeric") from exc
    if not math.isfinite(result):
        raise InputError(f"baseline workbook field '{field}' must be finite")
    return result


def load_baseline_xlsx(baseline_workbook: str | Path) -> Baseline:
    """Load and validate a frozen baseline workbook without reading raw history."""
    workbook = openpyxl.load_workbook(baseline_workbook, read_only=True, data_only=True)
    required_sheets = ("Metadata", "Thresholds", "Distributions", "Validation")
    missing_sheets = [sheet for sheet in required_sheets if sheet not in workbook.sheetnames]
    if missing_sheets:
        raise InputError(f"baseline workbook is missing sheets: {', '.join(missing_sheets)}")

    metadata_sheet = workbook["Metadata"]
    metadata = {
        str(row[0]): row[1]
        for row in metadata_sheet.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    }
    if str(_metadata_value(metadata, "baseline_schema_version")) != BASELINE_SCHEMA_VERSION:
        raise InputError("baseline workbook schema version is not supported by this calculator")
    if str(_metadata_value(metadata, "method_version")) != METHOD_VERSION:
        raise InputError("baseline workbook method version does not match this calculator")
    station_elevation_m = _as_float(_metadata_value(metadata, "station_elevation_m"), "station_elevation_m")
    altitude_factor = _as_float(_metadata_value(metadata, "uvp_altitude_factor"), "uvp_altitude_factor")
    if not math.isclose(station_elevation_m, STATION_ELEVATION_M, abs_tol=1e-9):
        raise InputError("baseline workbook station elevation does not match the current method")
    if not math.isclose(altitude_factor, uvp_altitude_factor(), abs_tol=1e-9):
        raise InputError("baseline workbook UVP altitude factor does not match the current method")
    reference_start = str(_metadata_value(metadata, "reference_start"))
    reference_end = str(_metadata_value(metadata, "reference_end"))
    _parse_date(reference_start)
    _parse_date(reference_end)

    threshold_values: dict[str, dict[float, float]] = {"uvp": {}, "clothing": {}}
    for index_name, quantile, value, *_ in workbook["Thresholds"].iter_rows(min_row=2, values_only=True):
        if index_name is None:
            continue
        name = str(index_name)
        if name not in threshold_values:
            raise InputError(f"baseline workbook has unsupported threshold index '{name}'")
        threshold_values[name][_as_float(quantile, f"{name}.quantile")] = _as_float(value, f"{name}.threshold_value")
    expected_quantiles = (0.2, 0.4, 0.6, 0.8)
    for name, values in threshold_values.items():
        if any(quantile not in values for quantile in expected_quantiles):
            raise InputError(f"baseline workbook is missing one or more {name} thresholds")

    min_temp_values: list[float] = []
    k_values: list[float] = []
    for _, min_temp, k_value, *_ in workbook["Distributions"].iter_rows(min_row=2, values_only=True):
        if min_temp is not None:
            min_temp_values.append(_as_float(min_temp, "min_temp_c_sorted"))
        if k_value is not None:
            k_values.append(_as_float(k_value, "wind_effect_k_sorted"))
    if not min_temp_values or not k_values or min_temp_values != sorted(min_temp_values) or k_values != sorted(k_values):
        raise InputError("baseline workbook distributions must be non-empty and sorted ascending")

    valid_counts = {
        str(metric): int(count)
        for metric, count, *_ in workbook["Validation"].iter_rows(min_row=2, values_only=True)
        if metric is not None and count is not None
    }
    required_counts = ("min_temp", "wind_effect_k", "uvp", "clothing")
    if any(metric not in valid_counts or valid_counts[metric] <= 0 for metric in required_counts):
        raise InputError("baseline workbook validation counts are incomplete")
    if valid_counts["min_temp"] != len(min_temp_values) or valid_counts["wind_effect_k"] != len(k_values):
        raise InputError("baseline workbook distributions do not match validation counts")
    workbook.close()
    return Baseline(
        min_temp_values=tuple(min_temp_values),
        k_values=tuple(k_values),
        uvp_thresholds=tuple(threshold_values["uvp"][quantile] for quantile in expected_quantiles),
        clothing_thresholds=tuple(threshold_values["clothing"][quantile] for quantile in expected_quantiles),
        reference_start=reference_start,
        reference_end=reference_end,
        valid_counts=valid_counts,
        station_elevation_m=station_elevation_m,
        uvp_altitude_factor=altitude_factor,
        source_workbook_name=str(_metadata_value(metadata, "source_workbook_name")),
        source_workbook_sha256=str(_metadata_value(metadata, "source_workbook_sha256")),
    )


def calculate_indices(forecast: dict[str, Any], baseline: Baseline) -> dict[str, Any]:
    """Return the five agreed daily indices for one forecast record."""
    target_date, t, tmin, rh, wind, sunshine = _validate_forecast(forecast)
    thi = calculate_thi(t, rh)
    thi_score, thi_level = thi_score_and_level(thi)
    k = calculate_k(t, wind, sunshine)
    k_score, k_level = k_score_and_level(k)
    uvp_unadjusted, uvp = calculate_uvp(target_date, sunshine)
    uvp_level = _five_level(uvp, baseline.uvp_thresholds, ("低暴露潜势", "较低暴露潜势", "中等暴露潜势", "较高暴露潜势", "高暴露潜势"))
    min_temp_rank, k_rank = _cold_rank(baseline.min_temp_values, tmin), _cold_rank(baseline.k_values, k)
    clothing = max(min_temp_rank, k_rank)
    clothing_level = _five_level(clothing, baseline.clothing_thresholds, ("极低保温需求", "较低保温需求", "中等保温需求", "较高保温需求", "高保温需求"))
    deviation = max(abs(thi_score), abs(k_score))
    comfort = 100.0 - 12.5 * deviation
    comfort_level = {100.0: "舒适", 75.0: "舒适", 50.0: "基本舒适", 25.0: "较不舒适", 0.0: "不舒适"}[comfort]
    return {
        "status": "ok",
        "input": {"date": target_date.isoformat(), "avg_temp_c": t, "min_temp_c": tmin, "avg_rh_percent": rh, "avg_wind_mps": wind, "sunshine_hours": round(sunshine, 4)},
        "indices": {
            "uvp": {
                "name_zh": "茶卡盐湖日照紫外暴露潜势指数",
                "value": round(uvp, 2),
                "unadjusted_value": round(uvp_unadjusted, 2),
                "altitude_factor": round(baseline.uvp_altitude_factor, 4),
                "station_elevation_m": baseline.station_elevation_m,
                "level": uvp_level,
                "note": "已按海拔每1000米约增加10%的近似规则调整；该指标不是标准UVI。",
            },
            "clothing": {"name_zh": "茶卡盐湖着装保温需求指数", "value": round(clothing, 2), "level": clothing_level, "min_temp_cold_rank": round(min_temp_rank, 2), "wind_effect_cold_rank": round(k_rank, 2)},
            "comfort": {"name_zh": "人体舒适度", "value": round(comfort, 2), "level": comfort_level, "max_deviation": deviation, "note": "表示未通过着装调整的原始天气舒适度。"},
            "thi": {"name_zh": "温湿度指数", "value": round(thi, 2), "score": thi_score, "level": thi_level},
            "wind_effect_k": {"name_zh": "风效指数", "value": round(k, 2), "score": k_score, "level": k_level},
        },
        "baseline": {
            "method_version": METHOD_VERSION,
            "reference_period": f"{baseline.reference_start} to {baseline.reference_end}", "valid_counts": baseline.valid_counts,
            "uvp_thresholds": [round(value, 3) for value in baseline.uvp_thresholds],
            "clothing_thresholds": [round(value, 3) for value in baseline.clothing_thresholds],
            "station_elevation_m": baseline.station_elevation_m,
            "uvp_altitude_factor": round(baseline.uvp_altitude_factor, 4),
            "source_workbook_name": baseline.source_workbook_name,
            "source_workbook_sha256": baseline.source_workbook_sha256,
        },
    }


def read_forecast_json(path: str | Path) -> dict[str, Any]:
    with open(path, "r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise InputError("input JSON must be a single object")
    return data
